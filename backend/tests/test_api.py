"""
The API, end to end against an in-memory database.

The isolation tests are the point of this file. Everything else here would be
caught by a careful read; "company B can fetch company A's recording" would
not, and it is the failure that matters most in a product where each customer's
calls are their own.
"""
from __future__ import annotations

import hashlib
import hmac
import json

import pytest


ADMIN = {"X-Admin-Key": "test-admin-key"}


# ---------------------------------------------------------------------------
# Health
# ---------------------------------------------------------------------------


def test_health_reports_capabilities(client):
    body = client.get("/api/health").json()
    assert body["status"] == "ok"
    assert set(body["capabilities"]) >= {
        "database", "credentialEncryption", "transcoding", "googleOAuth", "adminApi"
    }


def test_health_is_also_at_the_root(client):
    """A load balancer is configured once and should not know about /api."""
    assert client.get("/health").status_code == 200


# ---------------------------------------------------------------------------
# Onboarding
# ---------------------------------------------------------------------------


def test_create_company_returns_the_key_once(client):
    response = client.post(
        "/api/companies",
        headers=ADMIN,
        json={"phoneNumberId": "100", "name": "Acme", "wabaId": "w1", "accessToken": "tok"},
    )
    assert response.status_code == 201
    created = response.json()
    assert created["apiKey"].startswith("ca_")

    # And never again: no route hands a key back.
    again = client.get("/api/companies/100", headers=ADMIN).json()
    assert "apiKey" not in again


def test_duplicate_phone_number_id_is_a_conflict(client):
    payload = {"phoneNumberId": "100", "name": "Acme"}
    assert client.post("/api/companies", headers=ADMIN, json=payload).status_code == 201
    second = client.post("/api/companies", headers=ADMIN, json=payload)
    assert second.status_code == 409


def test_admin_routes_need_the_admin_key(client):
    assert client.get("/api/companies").status_code == 401
    assert client.get("/api/companies", headers={"X-Admin-Key": "wrong"}).status_code == 401


def test_admin_api_is_disabled_without_a_configured_key(client, monkeypatch):
    """A blank ADMIN_API_KEY must close the door, not leave it open."""
    from app.config import settings

    monkeypatch.setattr(settings, "admin_api_key", "")
    assert client.get("/api/companies", headers=ADMIN).status_code == 403


def test_stored_credentials_are_sealed(client, fake_db):
    client.post(
        "/api/companies",
        headers=ADMIN,
        json={"phoneNumberId": "100", "name": "Acme", "accessToken": "EAAG-secret"},
    )
    row = fake_db.tables["voice_tenants"]["100"]
    assert row["data"]["accessToken"].startswith("v1.")
    assert "EAAG-secret" not in json.dumps(row)


def test_patch_leaves_unmentioned_fields_alone(client, tenant_factory):
    tenant, _ = tenant_factory("200", "Before")
    client.patch("/api/companies/200", headers=ADMIN, json={"name": "After"})
    body = client.get("/api/companies/200", headers=ADMIN).json()
    assert body["name"] == "After"
    # The access token was never mentioned, so it is still there.
    assert body["configured"] is True


def test_rotate_key_invalidates_the_old_one(client, tenant_factory, auth):
    _, old_key = tenant_factory("200")
    assert client.get("/api/companies/me", headers=auth(old_key)).status_code == 200

    new_key = client.post("/api/companies/200/rotate-key", headers=ADMIN).json()["apiKey"]
    assert new_key != old_key
    assert client.get("/api/companies/me", headers=auth(old_key)).status_code == 401
    assert client.get("/api/companies/me", headers=auth(new_key)).status_code == 200


# ---------------------------------------------------------------------------
# Company authentication
# ---------------------------------------------------------------------------


def test_company_reads_itself(client, tenant_factory, auth):
    _, key = tenant_factory("300", "Harbour Clinics")
    body = client.get("/api/companies/me", headers=auth(key)).json()
    assert body["name"] == "Harbour Clinics"
    assert "accessToken" not in json.dumps(body)


def test_no_key_is_unauthorised(client):
    assert client.get("/api/calls").status_code == 401


def test_unknown_key_is_unauthorised(client, tenant_factory, auth):
    tenant_factory("300")
    assert client.get("/api/calls", headers=auth("ca_not-a-real-key")).status_code == 401


# ---------------------------------------------------------------------------
# Tenant isolation
# ---------------------------------------------------------------------------


@pytest.fixture
def two_companies(client, tenant_factory, auth):
    import asyncio

    from app.models.call import Call, Channel, RecordingState
    from app.repositories import calls as call_repo

    a, key_a = tenant_factory("100", "Alpha")
    b, key_b = tenant_factory("200", "Beta")

    call = Call(
        id="call-of-alpha",
        tenantId="100",
        channel=Channel.PHONE,
        counterparty="+923001112222",
        recordingPath="sb://tenants/100/calls/call-of-alpha.mp3",
        recordingMime="audio/mpeg",
        recordingState=RecordingState.READY,
    )
    asyncio.run(call_repo.save_call(call))
    return {"a": (a, key_a), "b": (b, key_b), "call": call}


def test_a_company_sees_only_its_own_calls(client, two_companies, auth):
    _, key_a = two_companies["a"]
    _, key_b = two_companies["b"]

    mine = client.get("/api/calls", headers=auth(key_a)).json()["calls"]
    assert [c["id"] for c in mine] == ["call-of-alpha"]

    theirs = client.get("/api/calls", headers=auth(key_b)).json()["calls"]
    assert theirs == []


def test_another_companys_call_is_not_found(client, two_companies, auth):
    _, key_b = two_companies["b"]
    # 404, not 403: confirming the id exists would itself leak something.
    assert client.get("/api/calls/call-of-alpha", headers=auth(key_b)).status_code == 404


def test_another_companys_recording_is_not_reachable(client, two_companies, auth):
    _, key_b = two_companies["b"]
    response = client.get("/api/calls/call-of-alpha/recording", headers=auth(key_b))
    assert response.status_code == 404


def test_another_companys_call_cannot_be_deleted(client, two_companies, auth):
    _, key_b = two_companies["b"]
    _, key_a = two_companies["a"]
    assert client.delete("/api/calls/call-of-alpha", headers=auth(key_b)).status_code == 404
    # Still there for its owner.
    assert client.get("/api/calls/call-of-alpha", headers=auth(key_a)).status_code == 200


def test_admin_may_look_at_a_named_company(client, two_companies):
    """The platform dashboard, without holding the customer's key."""
    response = client.get(
        "/api/calls", headers={**ADMIN, "X-Company-Id": "100"}
    )
    assert [c["id"] for c in response.json()["calls"]] == ["call-of-alpha"]


# ---------------------------------------------------------------------------
# Recordings
# ---------------------------------------------------------------------------


def test_recording_upload_stores_and_serves(client, tenant_factory, auth, monkeypatch):
    import asyncio

    from app.models.call import Call, Channel
    from app.repositories import calls as call_repo
    from app.services import storage

    tenant, key = tenant_factory("400")
    asyncio.run(
        call_repo.save_call(
            Call(id="c-1", tenantId="400", channel=Channel.BROWSER, counterparty="+923001112222")
        )
    )

    kept: dict[str, bytes] = {}

    async def fake_put(_tenant, call_id, data, mime, extension):
        kept[call_id] = data
        return storage.StoredRecording(f"sb://calls/{call_id}.{extension}", mime, len(data))

    async def fake_get(_tenant, reference):
        return kept.get(reference.split("/")[-1].split(".")[0])

    monkeypatch.setattr(storage, "put", fake_put)
    monkeypatch.setattr(storage, "get", fake_get)
    # Transcoding is exercised separately; here it would just need ffmpeg.
    monkeypatch.setattr("app.services.recordings.settings.transcode_recordings", False)

    upload = client.post(
        "/api/calls/c-1/recording",
        headers={**auth(key), "Content-Type": "audio/webm"},
        content=b"RIFFfake-audio-bytes",
    )
    assert upload.status_code == 201, upload.text
    assert upload.json()["bytes"] == len(b"RIFFfake-audio-bytes")

    played = client.get("/api/calls/c-1/recording", headers=auth(key))
    assert played.status_code == 200
    assert played.content == b"RIFFfake-audio-bytes"
    assert played.headers["accept-ranges"] == "bytes"


def test_recording_upload_rejects_non_audio(client, tenant_factory, auth):
    import asyncio

    from app.models.call import Call
    from app.repositories import calls as call_repo

    _, key = tenant_factory("400")
    asyncio.run(call_repo.save_call(Call(id="c-2", tenantId="400")))

    response = client.post(
        "/api/calls/c-2/recording",
        headers={**auth(key), "Content-Type": "application/pdf"},
        content=b"%PDF-1.4",
    )
    assert response.status_code == 415


def test_recording_upload_rejects_an_empty_body(client, tenant_factory, auth):
    import asyncio

    from app.models.call import Call
    from app.repositories import calls as call_repo

    _, key = tenant_factory("400")
    asyncio.run(call_repo.save_call(Call(id="c-3", tenantId="400")))

    response = client.post(
        "/api/calls/c-3/recording",
        headers={**auth(key), "Content-Type": "audio/webm"},
        content=b"",
    )
    assert response.status_code == 400


def test_recording_upload_closes_after_the_window(client, tenant_factory, auth):
    import asyncio

    from app.models.call import Call
    from app.repositories import calls as call_repo

    _, key = tenant_factory("400")
    asyncio.run(call_repo.save_call(Call(id="c-4", tenantId="400", startedAt=0.0)))

    response = client.post(
        "/api/calls/c-4/recording",
        headers={**auth(key), "Content-Type": "audio/webm"},
        content=b"late",
    )
    assert response.status_code == 409


def test_browser_can_report_why_it_failed(client, tenant_factory, auth):
    import asyncio

    from app.models.call import Call, RecordingState
    from app.repositories import calls as call_repo

    _, key = tenant_factory("400")
    asyncio.run(call_repo.save_call(Call(id="c-5", tenantId="400")))

    client.post(
        "/api/calls/c-5/recording/problem",
        headers=auth(key),
        json={"ok": False, "reason": "microphone blocked"},
    )
    call = asyncio.run(call_repo.get_call("400", "c-5"))
    assert call.recording_state == RecordingState.FAILED
    assert call.recording_error == "microphone blocked"


# ---------------------------------------------------------------------------
# Webhooks
# ---------------------------------------------------------------------------


def _signed(body: dict, secret: str) -> tuple[bytes, dict]:
    raw = json.dumps(body).encode()
    digest = hmac.new(secret.encode(), raw, hashlib.sha256).hexdigest()
    return raw, {"X-Hub-Signature-256": f"sha256={digest}", "Content-Type": "application/json"}


def _message_event(phone_number_id: str) -> dict:
    return {
        "entry": [
            {
                "changes": [
                    {
                        "field": "messages",
                        "value": {
                            "metadata": {"phone_number_id": phone_number_id},
                            "messages": [
                                {
                                    "id": "wamid.1",
                                    "from": "923001112222",
                                    "type": "text",
                                    "timestamp": "1700000000",
                                    "text": {"body": "when do admissions close?"},
                                }
                            ],
                        },
                    }
                ]
            }
        ]
    }


def test_signed_webhook_is_recorded(client, tenant_factory, auth):
    _, key = tenant_factory("500", appSecret="meta-secret")
    raw, headers = _signed(_message_event("500"), "meta-secret")

    response = client.post("/api/webhooks/whatsapp", content=raw, headers=headers)
    assert response.status_code == 200

    messages = client.get("/api/messages", headers=auth(key)).json()["messages"]
    assert messages[0]["body"] == "when do admissions close?"
    assert messages[0]["direction"] == "INBOUND"


def test_unsigned_webhook_is_refused(client, tenant_factory, auth):
    _, key = tenant_factory("500", appSecret="meta-secret")
    raw = json.dumps(_message_event("500")).encode()

    response = client.post(
        "/api/webhooks/whatsapp", content=raw, headers={"Content-Type": "application/json"}
    )
    assert response.status_code == 403
    assert client.get("/api/messages", headers=auth(key)).json()["messages"] == []


def test_webhook_is_verified_against_the_right_companys_secret(client, tenant_factory):
    """Company B's secret must not validate a delivery for company A."""
    tenant_factory("500", appSecret="secret-a")
    tenant_factory("600", appSecret="secret-b")

    raw, headers = _signed(_message_event("500"), "secret-b")
    assert client.post("/api/webhooks/whatsapp", content=raw, headers=headers).status_code == 403


def test_webhook_verification_challenge(client, tenant_factory):
    tenant_factory("500", verifyToken="my-verify-token")
    response = client.get(
        "/api/webhooks/whatsapp",
        params={
            "hub.mode": "subscribe",
            "hub.verify_token": "my-verify-token",
            "hub.challenge": "echo-me",
        },
    )
    assert response.status_code == 200
    assert response.text == "echo-me"


def test_webhook_verification_rejects_an_unknown_token(client, tenant_factory):
    tenant_factory("500", verifyToken="my-verify-token")
    response = client.get(
        "/api/webhooks/whatsapp",
        params={"hub.mode": "subscribe", "hub.verify_token": "guess", "hub.challenge": "x"},
    )
    assert response.status_code == 403


def test_inbound_call_event_is_logged(client, tenant_factory, auth):
    _, key = tenant_factory("500", appSecret="meta-secret")
    event = {
        "entry": [
            {
                "changes": [
                    {
                        "field": "calls",
                        "value": {
                            "metadata": {"phone_number_id": "500"},
                            "calls": [
                                {
                                    "id": "call-meta-1",
                                    "from": "923001112222",
                                    "event": "connect",
                                    "timestamp": "1700000000",
                                }
                            ],
                        },
                    }
                ]
            }
        ]
    }
    raw, headers = _signed(event, "meta-secret")
    client.post("/api/webhooks/whatsapp", content=raw, headers=headers)

    calls = client.get("/api/calls", headers=auth(key)).json()["calls"]
    assert calls[0]["channel"] == "WHATSAPP_CALL"
    assert calls[0]["status"] == "IN_PROGRESS"
    assert calls[0]["direction"] == "INBOUND"


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------


def test_workbook_downloads_and_opens(client, tenant_factory, auth):
    import asyncio
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models.call import Call, CallStatus
    from app.repositories import calls as call_repo

    _, key = tenant_factory("700", "Northwind")
    asyncio.run(
        call_repo.save_call(
            Call(
                id="x-1",
                tenantId="700",
                counterparty="+923001112222",
                status=CallStatus.COMPLETED,
                durationSeconds=214,
                answeredAt=1.0,
                summary="Asked about a delayed delivery.",
                transcript="Hello…",
            )
        )
    )

    response = client.get("/api/exports/calls.xlsx", headers=auth(key))
    assert response.status_code == 200
    assert "Northwind" in response.headers["content-disposition"]

    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Overview", "Calls", "Transcripts"]
    rows = list(book["Calls"].values)
    assert rows[0][0] == "Call ID"
    assert rows[1][0] == "x-1"
    assert rows[1][7] == "3:34"  # 214 seconds, as a length someone can read


def test_workbook_is_scoped_to_the_company(client, two_companies, auth):
    from io import BytesIO

    from openpyxl import load_workbook

    _, key_b = two_companies["b"]
    response = client.get("/api/exports/calls.xlsx", headers=auth(key_b))
    book = load_workbook(BytesIO(response.content))
    ids = [row[0] for row in list(book["Calls"].values)[1:]]
    assert "call-of-alpha" not in ids


# ---------------------------------------------------------------------------
# Google
# ---------------------------------------------------------------------------


def test_drive_status_before_connecting(client, tenant_factory, auth):
    _, key = tenant_factory("800")
    body = client.get("/api/google/status", headers=auth(key)).json()
    assert body["connected"] is False


def test_connect_returns_a_google_url(client, tenant_factory, auth, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "google_client_id", "client-id")
    monkeypatch.setattr(settings, "google_client_secret", "client-secret")
    monkeypatch.setattr(settings, "public_base_url", "https://api.example.com")

    _, key = tenant_factory("800")
    url = client.post("/api/google/connect", headers=auth(key)).json()["url"]
    assert url.startswith("https://accounts.google.com/o/oauth2/v2/auth?")
    # Without these two, Google returns no refresh token and the connection
    # silently dies after an hour.
    assert "access_type=offline" in url
    assert "prompt=consent" in url


def test_sync_report_refuses_when_drive_is_not_connected(client, tenant_factory, auth):
    _, key = tenant_factory("800")
    assert client.post("/api/google/sync-report", headers=auth(key)).status_code == 409


# ---------------------------------------------------------------------------
# Error shape
# ---------------------------------------------------------------------------


def test_errors_share_one_envelope(client):
    body = client.get("/api/calls").json()
    assert set(body["error"]) >= {"message", "code"}


def test_validation_errors_name_the_field(client, tenant_factory, auth):
    _, key = tenant_factory("900")
    body = client.post("/api/messages/text", headers=auth(key), json={"to": "+92300"}).json()
    assert body["error"]["code"] == "validation_error"
    assert any("body" in d["field"] for d in body["error"]["details"])
